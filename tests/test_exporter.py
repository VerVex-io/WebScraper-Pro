import csv
import io
import json

import pytest
from openpyxl import load_workbook
from PIL import Image

from tests.conftest import FakeResponse
from webscraper_pro import exporter
from webscraper_pro.security import SecurityError

ROWS = [
    {
        "name": '=HYPERLINK("http://evil.test","x")',
        "price": "+1",
        "url": "https://a.test/1",
        "image_url": "N/A",
    },
    {"name": "Normal", "price": "$5", "url": "@SUM(1)", "image_url": "N/A"},
]


def test_csv_neutralises_formulas(tmp_path):
    p = tmp_path / "o.csv"
    exporter.export_csv(ROWS, p)
    rows = list(csv.reader(p.open(encoding="utf-8-sig")))
    assert rows[1][0].startswith("'=") and rows[1][1] == "'+1" and rows[2][2] == "'@SUM(1)"


def test_xlsx_neutralises_formulas(tmp_path):
    p = tmp_path / "o.xlsx"
    exporter.export_xlsx(ROWS, p, images=False)
    ws = load_workbook(p).active
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            assert cell.data_type != "f", f"formula survived in {cell.coordinate}"
    assert ws["A2"].value.startswith("'=")


def test_json_is_raw_data(tmp_path):
    p = tmp_path / "o.json"
    exporter.export_json(ROWS, p)
    assert json.loads(p.read_text())[0]["name"].startswith("=")  # JSON stays faithful


def test_output_path_is_contained(tmp_path):
    p = exporter.output_path(tmp_path / "out", "../../site", "../../../etc/passwd", "csv")
    assert p.parent == (tmp_path / "out").resolve()
    assert ".." not in p.name and "/" not in p.name


def _png(size=(10, 10)) -> bytes:
    buf = io.BytesIO()
    Image.new("RGB", size, "red").save(buf, "PNG")
    return buf.getvalue()


def test_thumbnail_ok(monkeypatch):
    monkeypatch.setattr(exporter, "safe_get", lambda *a, **k: (_png((800, 600)), FakeResponse()))
    out = Image.open(io.BytesIO(exporter.fetch_thumbnail("https://a.test/i.png")))
    assert max(out.size) <= 200


def test_thumbnail_rejects_non_image(monkeypatch):
    monkeypatch.setattr(exporter, "safe_get", lambda *a, **k: (b"<html>not an image</html>", FakeResponse()))
    with pytest.raises(Exception):
        exporter.fetch_thumbnail("https://a.test/i.png")


def test_thumbnail_rejects_decompression_bomb(monkeypatch):
    # 10k x 10k = 100M pixels, but compresses to a few KB.
    buf = io.BytesIO()
    Image.new("1", (10000, 10000)).save(buf, "PNG")
    assert len(buf.getvalue()) < 1_000_000
    monkeypatch.setattr(exporter, "safe_get", lambda *a, **k: (buf.getvalue(), FakeResponse()))
    with pytest.raises(Image.DecompressionBombError):
        exporter.fetch_thumbnail("https://a.test/bomb.png")


def test_xlsx_embeds_images_and_survives_failures(tmp_path, monkeypatch):
    calls = {"n": 0}

    def fake_thumb(url, **k):
        calls["n"] += 1
        if "bad" in url:
            raise SecurityError("blocked")
        return _png()

    monkeypatch.setattr(exporter, "fetch_thumbnail", fake_thumb)
    rows = [
        {"name": "a", "price": "1", "url": "u", "image_url": "https://a.test/ok.png"},
        {"name": "b", "price": "2", "url": "u", "image_url": "https://a.test/bad.png"},
        {"name": "c", "price": "3", "url": "u", "image_url": "https://a.test/ok2.png"},
    ]
    logs = []
    p = tmp_path / "o.xlsx"
    exporter.export_xlsx(rows, p, max_images=2, log=logs.append)
    assert len(load_workbook(p).active._images) == 2
    assert any("skipped" in m for m in logs)


def test_export_all_unknown_format(tmp_path):
    with pytest.raises(ValueError):
        exporter.export_all(ROWS, tmp_path, "s", "k", formats=("exe",))


def test_export_all_writes_three_files(tmp_path):
    files = exporter.export_all(ROWS, tmp_path, "s", "k", images=False)
    assert sorted(f.suffix for f in files) == [".csv", ".json", ".xlsx"]
